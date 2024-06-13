# passo-a-passo
# -passar info da planilha para .word
# -salvar .word em uma pasta especifica (contratos)
# -repetir para o resto da planilha

from openpyxl import load_workbook
from docx import Document
from datetime import datetime

#entrar na planilha
workbook = load_workbook('fornecedores.xlsx')
sheet_products = workbook['Produtos']

for line in sheet_products.iter_rols(min_row=2,values_only=True):
    enterprise_name, address, city, state, zip_code, phone_number, email, sector = line #unpacking
    
    word_archive = Document()
    word_archive.add_heading('Contrato de Prestação de Serviço',0)
    
    contract_text = f"""
        Este contrato de prestação de serviços é feito entre {enterprise_name}, com endereço em {address}, 
    {city}, {state}, CEP {zip_code}, doravante denominado FORNECEDOR, e a empresa CONTRATANTE.

    Pelo presente instrumento particular, as partes têm, entre si, justo e acordado o seguinte:

    1. OBJETO DO CONTRATO
    O FORNECEDOR compromete-se a fornecer à CONTRATANTE os serviços/material de acordo com as especificações acordadas, respeitando os padrões de qualidade e os prazos estipulados.

    2. PRAZO
    Este contrato tem prazo de vigência de 12 (doze) meses, iniciando-se na data de sua assinatura, podendo ser renovado conforme acordo entre as partes.

    3. VALOR E FORMA DE PAGAMENTO
    O valor dos serviços prestados será acordado conforme as demandas da CONTRATANTE e a capacidade de entrega do FORNECEDOR. Os pagamentos serão realizados mensalmente, mediante apresentação de nota fiscal.

    4. CONFIDENCIALIDADE
    Todas as informações trocadas entre as partes durante a vigência deste contrato serão tratadas como confidenciais.

    Para firmeza e como prova de assim haverem justo e contratado, as partes assinam o presente contrato em duas vias de igual teor e forma.

    FORNECEDOR: {enterprise_name}
    E-mail: {email}

    CONTRATANTE: [NOME CONTRATANTE]
    E-mail: [E-MAIL CONTRATANTE]

    [CIDADE],{datetime.now().strftime('%d/%M/%Y')}
    """
    
    